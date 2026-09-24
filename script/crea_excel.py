"""Crea musei_link.xlsx a partire da DATI/verifica.json.

Uso: python3 crea_excel.py DATI FILE_EXCEL
"""
import json
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

DATI, OUT = sys.argv[1], sys.argv[2]

COLONNE = ["Regione", "Provincia", "Comune", "Nome", "Scheda Abbonamento Musei", "Sito ufficiale",
           "Fonte del link", "Esito", "Codice HTTP", "URL finale", "Note", "Data verifica"]
LARGHEZZE = [16, 10, 22, 44, 40, 40, 13, 20, 11, 40, 70, 13]
ORDINE_REGIONI = {"Lombardia": 0, "Piemonte": 1, "Valle d'Aosta": 2}
COLORI = {
    "FUNZIONA": "C6EFCE", "REINDIRIZZA": "DDEBF7", "GENERICO": "FFF2CC", "SOLO SOCIAL": "EDE0F7",
    "NON FUNZIONA": "F8CBAD", "BLOCCATO": "FCE4D6", "DUBBIO": "FFE699", "nessun sito trovato": "D9D9D9",
}
DA_CONTROLLARE = {"BLOCCATO", "DUBBIO", "NON FUNZIONA", "nessun sito trovato"}

FONT = Font(name="Arial", size=10)
FONT_LINK = Font(name="Arial", size=10, color="0563C1", underline="single")
FONT_TESTA = Font(name="Arial", size=10, bold=True, color="FFFFFF")
FILL_TESTA = PatternFill("solid", fgColor="305496")


def riga(m):
    return [m["regione"], m["provincia"], m["comune"], m["nome"], m["scheda"],
            m["sito"] or "nessun sito trovato", m.get("fonte") or "", m["esito"],
            m.get("codice"), m.get("finale") or "", "; ".join(m.get("note") or []), m.get("data", "")]


def foglio(ws, righe):
    ws.append(COLONNE)
    for c in ws[1]:
        c.font, c.fill = FONT_TESTA, FILL_TESTA
        c.alignment = Alignment(vertical="center", wrap_text=True)
    for r in righe:
        ws.append(r)
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font = FONT
            c.alignment = Alignment(vertical="top", wrap_text=(c.column in (4, 11)))
        for col in (5, 6, 10):
            c = row[col - 1]
            if isinstance(c.value, str) and c.value.startswith("http"):
                c.hyperlink = c.value
                c.font = FONT_LINK
        esito = row[7]
        if esito.value in COLORI:
            esito.fill = PatternFill("solid", fgColor=COLORI[esito.value])
    for i, w in enumerate(LARGHEZZE, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLONNE))}{ws.max_row}"


def main():
    musei = json.load(open(f"{DATI}/verifica.json", encoding="utf-8"))
    musei.sort(key=lambda m: (ORDINE_REGIONI.get(m["regione"], 9), m["provincia"], m["comune"].lower(),
                              m["nome"].lower()))
    wb = Workbook()
    ws = wb.active
    ws.title = "Musei"
    foglio(ws, [riga(m) for m in musei])
    ws2 = wb.create_sheet("Da controllare")
    foglio(ws2, [riga(m) for m in musei if m["esito"] in DA_CONTROLLARE])
    wb.save(OUT)
    print("righe:", len(musei), "| da controllare:", ws2.max_row - 1)


if __name__ == "__main__":
    main()
